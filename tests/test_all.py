"""
單元測試：python -m unittest discover -s tests -v

真實影片的回歸測試只在設定環境變數 TW_SAMPLES（放樣本影片的資料夾）時執行；
樣本影片是使用者的監測資料，不可以放進 repository。
"""
import datetime as dt
import glob
import os
import sys
import tempfile
import unittest

import numpy as np

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from trainwatch import core, export, osd, selftest  # noqa: E402

_TMP = []


def _mkdtemp():
    """測試用的暫存資料夾，全部測試跑完就刪掉（假影片很大，不刪會把磁碟塞滿）"""
    d = tempfile.mkdtemp(prefix='tw_test_')
    _TMP.append(d)
    return d


def tearDownModule():
    import shutil
    for d in _TMP:
        shutil.rmtree(d, ignore_errors=True)


class TestOsdFormat(unittest.TestCase):
    def test_tokens_and_roundtrip(self):
        fmt = 'YYYY-MM-DD HH:MM:SS'
        d = dt.datetime(2026, 1, 2, 23, 4, 7)
        self.assertEqual(osd.render(d, fmt), '2026-01-02 23:04:07')
        self.assertEqual(osd.parse('2026-01-02 23:04:07', fmt), d)

    def test_other_format(self):
        fmt = 'DD/MM/YY HH:MM:SS'
        d = dt.datetime(2026, 3, 5, 7, 8, 9)
        self.assertEqual(osd.render(d, fmt), '05/03/26 07:08:09')
        self.assertEqual(osd.parse('05/03/26 07:08:09', fmt), d)

    def test_invalid(self):
        self.assertIsNone(osd.parse('2026-13-02 23:04:07', 'YYYY-MM-DD HH:MM:SS'))
        self.assertIsNone(osd.parse('2026-01-02 25:04:07', 'YYYY-MM-DD HH:MM:SS'))
        with self.assertRaises(ValueError):
            osd.tokens('HH:MM:SS')

    def test_read_rendered_digits(self):
        """用內建字樣拼出一張字幕圖，應該讀回同一個字串"""
        import cv2
        fmt = osd.DEFAULT_FORMAT
        T = osd.default_templates()
        text = '2026-01-02 13:58:59'
        x0, y0 = 14, 8
        img = np.zeros((40, 260), np.uint8)
        for i, ch in enumerate(text):
            if ch.isdigit():
                cell = (T[int(ch)] * 80 + 150).clip(0, 255).astype(np.uint8)
                img[y0:y0 + 16, x0 + i * 10:x0 + i * 10 + 10] = cell
        frame = cv2.cvtColor(img, cv2.COLOR_GRAY2BGR)
        d, txt, w = osd.read_datetime(frame, osd.DEFAULT_RECT, fmt, T)
        self.assertEqual(txt, text)
        self.assertEqual(d, dt.datetime(2026, 1, 2, 13, 58, 59))


class TestCore(unittest.TestCase):
    def test_number_events(self):
        evs = [dict(start=3, valid=True), dict(start=1, valid=False), dict(start=2, valid=True)]
        core.number_events(evs)
        self.assertEqual([(e['start'], e['no']) for e in evs], [(1, 1), (2, 1), (3, 2)])

    def test_fmt_time(self):
        ts = dt.datetime(2026, 1, 2, 7, 6, 5).timestamp() + 0.35
        self.assertEqual(core.fmt_time(ts), '07:06:05.3')
        self.assertEqual(core.fmt_time(ts, True), '2026-01-02 07:06:05.3')

    def test_profile_json(self):
        p = core.Profile(name='站A')
        q = core.Profile.from_dict(__import__('json').loads(p.to_json()))
        self.assertEqual(q.name, '站A')
        self.assertEqual(q.osd_rect, p.osd_rect)

    def test_selftest_video(self):
        """假列車影片：兩筆、時間誤差 < 0.4 秒、方向正確、有截圖與短片"""
        r = selftest.run()
        self.assertTrue(r['ok'], r['problems'])


class TestScenarios(unittest.TestCase):
    """驗收測試（假影片）：列車停駛、兩列車重疊、車廂空隙、曝光變化。允許誤差：車頭 ±0.3 秒、車尾 ±0.4 秒"""

    def run_case(self, seconds, **kw):
        d = _mkdtemp()
        got, exp = selftest.run_scene(os.path.join(d, 's.avi'), seconds, **kw)
        return got, exp

    def check(self, got, want, dirs=None):
        self.assertEqual(len(got), len(want), got)
        for i, (g, (a, b)) in enumerate(zip(got, want)):
            self.assertLess(abs(g[0] - a), 0.3, (g, a, b))
            self.assertLess(abs(g[1] - b), 0.4, (g, a, b))
            if dirs:
                self.assertEqual(g[2], dirs[i], g)

    def test1_normal(self):
        got, exp = self.run_case(20, trains=[dict(t0=5, d=1)])
        self.check(got, exp, ['往右'])

    def test2_stop_5s(self):
        got, exp = self.run_case(30, trains=[dict(t0=5, d=1, stops=[(6.0, 5)])])
        self.check(got, exp, ['往右'])
        self.assertIn('停止', got[0][3])

    def test3_stop_35s(self):
        got, exp = self.run_case(120, trains=[dict(t0=5, d=-1, stops=[(6.0, 35)])])
        self.check(got, exp, ['往左'])
        self.assertGreater(got[0][1] - got[0][0], 35)

    def test4_two_trains_overlap(self):
        """A 10～30 秒、B 20～40 秒（雙軌、反方向）→ 一筆 10～40 秒"""
        got, exp = self.run_case(50, trains=[dict(t0=10, d=1, speed=40, length=800, y=(166, 188)),
                                             dict(t0=20, d=-1, speed=40, length=800, y=(192, 214))])
        self.check(got, [(10, 40)])

    def test5_two_trains_clear_gap(self):
        got, exp = self.run_case(50, trains=[dict(t0=10, d=1, speed=40, length=400),
                                             dict(t0=30, d=-1, speed=40, length=400)])
        self.check(got, exp, ['往右', '往左'])

    def test6_coach_gaps(self):
        got, exp = self.run_case(40, trains=[dict(t0=5, d=1, speed=60, length=1200, car=180, gap=30)])
        self.check(got, exp, ['往右'])

    def test7_exposure_no_train(self):
        got, exp = self.run_case(40, exposure=[(10, 20, 0.6), (25, 26, 1.4)], blobs=[(30, 33, 420, 190, 30)])
        self.assertEqual(got, [])

    def test7b_darken_after_train(self):
        got, exp = self.run_case(30, trains=[dict(t0=5, d=1)], exposure=[(7.3, 14, 0.75)])
        self.check(got, exp, ['往右'])

    def test8_local_sunlight_no_train(self):
        """白天只有軌道一帶被陽光照亮（全畫面亮度修正抓不到）→ 不可以變成列車"""
        got, exp = self.run_case(60, local_light=[(10, 40, 1.35)])
        self.assertEqual(got, [])

    def test8b_train_during_sunlight(self):
        """陽光照亮期間有列車經過 → 只記列車那一段（不可以從陽光照到的時候就開始）"""
        got, exp = self.run_case(60, trains=[dict(t0=25, d=-1)], local_light=[(10, 50, 1.35)])
        self.check(got, exp, ['往左'])

    def test8c_scene_is_daytime(self):
        """確認假影片被當成白天（彩度夠高），上面兩個測試才有意義"""
        import cv2
        d = _mkdtemp(); f = os.path.join(d, 's.avi')
        selftest.make_scene(f, 2)
        ok, fr = cv2.VideoCapture(f).read()
        sat = float(cv2.cvtColor(fr, cv2.COLOR_BGR2HSV)[:, :, 1].mean())
        self.assertGreater(sat, core.Profile().day_saturation)

    def test_shots_cancel(self):
        """產生截圖與短片時按取消，要真的停下來"""
        d = _mkdtemp()
        f = os.path.join(d, 's.avi')
        selftest.make_scene(f, 40, trains=[dict(t0=5, d=1), dict(t0=25, d=-1)])
        prof = core.Profile()
        evs = core.number_events(core.process([f], [0.0], prof))
        core.save_frames_and_clips(evs, prof, os.path.join(d, 'out'), cancel=lambda: True, files=[f])
        self.assertEqual(sum(len(e['shots']) for e in evs), 0)


class TestRound2(unittest.TestCase):
    """v1.0.9：GPT 第二輪檢查的問題（修正清單 #45～#62）。先寫測試確認問題存在，再修"""

    def scene(self, name, seconds, **kw):
        d = _mkdtemp()
        f = os.path.join(d, name)
        exp = selftest.make_scene(f, seconds, **kw)
        return d, f, exp

    def run_files(self, files, bases, prof=None, cancel=None):
        prof = prof or core.Profile()
        return core.number_events(core.process(files, bases, prof, cancel=cancel)), prof

    def valid(self, evs):
        return [e for e in evs if e['valid']]

    def test45_video_ends_with_train(self):
        """#45 影片 20 秒，列車第 15 秒到達、第 22 秒才會離開 → 車尾不可以當成已確認"""
        d, f, _ = self.scene('a.avi', 20, trains=[dict(t0=15, d=1, speed=40)])
        evs, _ = self.run_files([f], [0.0])
        v = self.valid(evs)
        self.assertEqual(len(v), 1, evs)
        self.assertTrue(v[0]['need_check'])
        self.assertFalse(v[0].get('end_known', True))
        self.assertIn('影片結束時列車仍在參考線上', v[0]['reason'])

    def test45b_cancel_during_train(self):
        """#45 判讀途中按取消、當時列車正在通過 → 車尾不可以當成已確認"""
        d, f, _ = self.scene('a.avi', 40, trains=[dict(t0=5, d=1, speed=40)])
        n = {'k': 0}
        def cancel():
            n['k'] += 1
            return n['k'] > 3          # 每 30 格檢查一次 → 約第 6～8 秒取消
        evs, _ = self.run_files([f], [0.0], cancel=cancel)
        v = self.valid(evs)
        self.assertEqual(len(v), 1, evs)
        self.assertTrue(v[0]['need_check'])
        self.assertFalse(v[0].get('end_known', True))
        self.assertIn('取消', v[0]['reason'])

    def test46_static_limit_end_unknown(self):
        """#46→#67 停住超過 max_static（測試用 15 秒）、影片結束時還沒開走 → 不提早結束（v1.0.10），
        到影片結束才收尾，標「車尾未確認」與「長時間沒有動靜」"""
        d, f, _ = self.scene('a.avi', 150, trains=[dict(t0=60, d=1, stops=[(61.0, 1000)])])
        prof = core.Profile(); prof.max_static = 15
        evs, _ = self.run_files([f], [0.0], prof)
        v = self.valid(evs)
        self.assertEqual(len(v), 1, evs)
        self.assertGreater(v[0]['end'], 145)
        self.assertTrue(v[0]['need_check'])
        self.assertFalse(v[0].get('end_known', True))
        self.assertIn('車尾未確認', v[0]['reason'])
        self.assertIn('長時間', v[0]['reason'])

    def test47_static_limit_across_files(self):
        """#47→#67 列車在 A 停住、B 整支還停著 → 一筆，從 A 到 B 結束（v1.0.10 起不再因 max_static 提早結束）；
        影片內秒數不可以是負的"""
        d = _mkdtemp()
        a, b = os.path.join(d, 'a.avi'), os.path.join(d, 'b.avi')
        selftest.make_scene(a, 150, trains=[dict(t0=60, d=1, stops=[(61.0, 1000)])])
        selftest.make_scene(b, 40, trains=[dict(t0=-90, d=1, stops=[(-89.0, 1000)])])
        prof = core.Profile(); prof.max_static = 100
        evs, _ = self.run_files([a, b], [0.0, 150.0], prof)
        v = self.valid(evs)
        self.assertEqual(len(v), 1, [(e['start'], e['end']) for e in v])
        self.assertEqual(v[0]['start_file'], a)
        self.assertEqual(v[0]['end_file'], b, v[0])
        self.assertGreaterEqual(v[0]['end_pos'], 0.0)
        self.assertFalse(v[0]['end_known'])
        self.assertIn('長時間', v[0]['reason'])

    def test49_gap_between_files(self):
        """#49 A 結尾列車還在、B 在 50 分鐘後開始也有列車 → 不可以接成一筆；A 那筆車尾未確認"""
        d = _mkdtemp()
        a, b = os.path.join(d, 'a.avi'), os.path.join(d, 'b.avi')
        selftest.make_scene(a, 20, trains=[dict(t0=15, d=1, speed=40)])
        selftest.make_scene(b, 20, trains=[dict(t0=-3, d=1, speed=40)])
        evs, _ = self.run_files([a, b], [0.0, 3000.0])
        v = self.valid(evs)
        self.assertEqual(len(v), 2, [(e['start'], e['end']) for e in v])
        self.assertLess(v[0]['end'], 25)
        self.assertFalse(v[0].get('end_known', True))
        self.assertIn('中斷', v[0]['reason'])
        self.assertGreaterEqual(v[1]['start'], 3000.0)
        self.assertTrue(v[1]['need_check'])

    def test51_vertical_sunlight(self):
        """#51 上下走向：陽光照亮（無車）→ 不可以有列車；陽光下有列車 → 時間正確"""
        d, f, _ = self.scene('a.avi', 60, vertical=True, local_light=[(10, 50, 1.35)])
        evs, _ = self.run_files([f], [0.0], selftest.vertical_profile())
        self.assertEqual(self.valid(evs), [])
        d, f, exp = self.scene('b.avi', 60, vertical=True, trains=[dict(t0=25, d=-1)], local_light=[(10, 50, 1.35)])
        evs, _ = self.run_files([f], [0.0], selftest.vertical_profile())
        v = self.valid(evs)
        self.assertEqual(len(v), 1, [(e['start'], e['end']) for e in v])
        self.assertLess(abs(v[0]['start'] - exp[0][0]), 0.3)
        self.assertLess(abs(v[0]['end'] - exp[0][1]), 0.4)
        self.assertEqual(v[0]['direction'], '往上')

    def test51b_vertical_band(self):
        """#51 紋理比對的範圍要沿軌道方向加寬（上下走向＝往上下加寬）"""
        p = selftest.vertical_profile()
        det = core.Detector(p, core.work_crop(p, 360, 640))
        ys, xs = np.nonzero(det.band)
        ry, rx = np.nonzero(det.ref_mask)
        self.assertGreaterEqual(ys.max() - ys.min(), ry.max() - ry.min() + 18)

    def test53_ir_switch_no_train(self):
        """#53 傍晚切到紅外線（無車）→ 不可以有列車"""
        d, f, _ = self.scene('a.avi', 60, ir_at=20)
        evs, _ = self.run_files([f], [0.0])
        self.assertEqual(self.valid(evs), [], [(e['start'], e['end'], e['reason']) for e in evs])

    def test53b_ir_switch_then_train(self):
        """#53 切到紅外線之後來一列車 → 要判讀到，時間正確"""
        d, f, exp = self.scene('a.avi', 60, ir_at=15, trains=[dict(t0=35, d=1)])
        evs, _ = self.run_files([f], [0.0])
        v = self.valid(evs)
        self.assertEqual(len(v), 1, [(e['start'], e['end'], e['reason']) for e in evs])
        self.assertLess(abs(v[0]['start'] - exp[0][0]), 0.5)
        self.assertLess(abs(v[0]['end'] - exp[0][1]), 0.6)

    def test56_train_standing_at_start(self):
        """#56 第一支影片一開頭列車就停在參考線上。
        停 10 秒：時間要對、標需確認（影片一開始就有車）。
        停 60、100 秒（超過開頭取樣的一半）：分不出列車是開頭停著還是後來停進來，不可以默默給時間，要標「背景不確定」"""
        for stop in (10, 60, 100):
            d, f, exp = self.scene('a.avi', 300, trains=[dict(t0=-1, d=1, stops=[(0.0, stop)])])
            evs, _ = self.run_files([f], [0.0])
            v = self.valid(evs)
            self.assertGreaterEqual(len(v), 1, (stop, [(e['start'], e['end'], e['valid'], e['reason']) for e in evs]))
            self.assertTrue(all(e['need_check'] for e in v), stop)
            if stop == 10:
                self.assertEqual(len(v), 1)
                self.assertLess(v[0]['start'], 1.0)
                self.assertLess(abs(v[0]['end'] - exp[0][1]), 0.6, (stop, v[0]['end'], exp))
            else:
                self.assertTrue(all('背景不確定' in e['reason'] for e in v), (stop, [e['reason'] for e in v]))

    def clip_seconds(self, path):
        import cv2
        c = cv2.VideoCapture(path)
        n, fps = c.get(cv2.CAP_PROP_FRAME_COUNT), c.get(cv2.CAP_PROP_FPS)
        c.release()
        return n / fps

    def test54_clip_three_files(self):
        """#54 列車停住、事件跨 A、B、C 三支影片 → 短片要包含中間的 B"""
        d = _mkdtemp()
        fs = [os.path.join(d, x + '.avi') for x in 'abc']
        for i, f in enumerate(fs):
            selftest.make_scene(f, 15, trains=[dict(t0=10 - 15 * i, d=1, stops=[(11.0 - 15 * i, 25)])])
        evs, prof = self.run_files(fs, [0.0, 15.0, 30.0])
        v = self.valid(evs)
        self.assertEqual(len(v), 1, [(e['start'], e['end']) for e in v])
        dur = v[0]['end'] - v[0]['start']
        out = os.path.join(d, 'out')
        core.save_frames_and_clips(evs, prof, out, files=fs, bases=[0.0, 15.0, 30.0])
        sec = self.clip_seconds(os.path.join(out, '短片', v[0]['clip']))
        self.assertLess(abs(sec - (dur + 6)), 1.0, (sec, dur))

    def test59_long_event_clip(self):
        """#59 超過 60 秒的事件：短片只有車頭前後、車尾前後兩段（各 3＋10 秒）"""
        d, f, exp = self.scene('a.avi', 200, trains=[dict(t0=60, d=1, stops=[(61.0, 80)])])
        evs, prof = self.run_files([f], [0.0])
        v = self.valid(evs)
        self.assertEqual(len(v), 1)
        self.assertGreater(v[0]['end'] - v[0]['start'], 60)
        out = os.path.join(d, 'out')
        core.save_frames_and_clips(evs, prof, out, files=[f], bases=[0.0])
        sec = self.clip_seconds(os.path.join(out, '短片', v[0]['clip']))
        self.assertLess(abs(sec - 26), 1.0, sec)
        self.assertEqual(len(v[0]['shots']), 3)

    def test60b_media_tag_no_overwrite(self):
        """#60 已排除改成列車後補做短片：用原本的檔名前綴（X001），不可以蓋到其他列車的截圖與短片"""
        d, f, _ = self.scene('a.avi', 40, trains=[dict(t0=5, d=1)], blobs=[(25, 28, 335, 190, 22)])
        evs, prof = self.run_files([f], [0.0])
        out = os.path.join(d, 'out')
        core.save_frames_and_clips(evs, prof, out, files=[f], bases=[0.0])
        x = next(e for e in evs if not e['valid'])
        t = next(e for e in evs if e['valid'])
        before = os.path.getmtime(os.path.join(out, '短片', t['clip']))
        x['valid'] = True
        core.number_events(evs)                   # 編號重排：x 變成第 2 筆
        core.save_frames_and_clips([x], prof, out, files=[f], bases=[0.0])
        self.assertEqual(x['clip'], 'X001.mp4')
        self.assertTrue(all(s.startswith('X001_') for s in x['shots']))
        self.assertEqual(os.path.getmtime(os.path.join(out, '短片', t['clip'])), before)

    def test60_excluded_no_clip_peak_shot(self):
        """#60 已排除的事件：不產生短片；截圖有一張是「變化最大」的那一刻（看得出排除原因）"""
        d, f, _ = self.scene('a.avi', 30, blobs=[(10, 13, 335, 190, 22)])
        evs, prof = self.run_files([f], [0.0])
        x = [e for e in evs if not e['valid']]
        self.assertGreaterEqual(len(x), 1, evs)
        out = os.path.join(d, 'out')
        core.save_frames_and_clips(evs, prof, out, files=[f], bases=[0.0])
        self.assertEqual(x[0]['clip'], '')
        self.assertTrue(any('變化最大' in s for s in x[0]['shots']), x[0]['shots'])
        self.assertTrue(all(os.path.exists(os.path.join(out, '截圖', s)) for s in x[0]['shots']))


class TestRound3(unittest.TestCase):
    """v1.0.10：GPT 第三輪檢查的問題（修正清單 #66～#70）。先寫測試確認問題存在，再修"""
    scene = TestRound2.scene
    run_files = TestRound2.run_files
    valid = TestRound2.valid

    def test66a_video_end_while_confirming(self):
        """#66 車尾看起來離開了，但還沒等滿確認時間影片就結束 → 車尾未確認"""
        d, f, exp = self.scene('a.avi', 20, trains=[dict(t0=17, d=1)])
        evs, _ = self.run_files([f], [0.0])
        v = self.valid(evs)
        self.assertEqual(len(v), 1, evs)
        self.assertLess(abs(v[0]['end'] - exp[0][1]), 0.4)          # 時間照樣保留最後看到的車尾
        self.assertFalse(v[0]['end_known']); self.assertTrue(v[0]['need_check'])
        self.assertIn('確認時間', v[0]['reason'])

    def test66b_cancel_while_confirming(self):
        """#66 按取消時正在等確認 → 車尾未確認"""
        d, f, exp = self.scene('a.avi', 30, trains=[dict(t0=13, d=1)])   # 車尾約 15.2 秒離開
        n = {'k': 0}
        def cancel():
            n['k'] += 1
            return n['k'] >= 8          # 每 30 格（2 秒）檢查一次 → 第 16 秒取消
        evs, _ = self.run_files([f], [0.0], cancel=cancel)
        v = self.valid(evs)
        self.assertEqual(len(v), 1, evs)
        self.assertFalse(v[0]['end_known']); self.assertTrue(v[0]['need_check'])

    def test66c_gap_while_confirming(self):
        """#66 錄影中斷時正在等確認 → 車尾未確認"""
        d = _mkdtemp()
        a, b = os.path.join(d, 'a.avi'), os.path.join(d, 'b.avi')
        selftest.make_scene(a, 20, trains=[dict(t0=17, d=1)])
        selftest.make_scene(b, 20)
        evs, _ = self.run_files([a, b], [0.0, 3000.0])
        v = self.valid(evs)
        self.assertEqual(len(v), 1, evs)
        self.assertFalse(v[0]['end_known']); self.assertTrue(v[0]['need_check'])

    def test66d_normal_end_still_known(self):
        """#66 正常離開、後面有足夠的無車畫面 → 車尾已確認（不可以改壞）"""
        d, f, exp = self.scene('a.avi', 20, trains=[dict(t0=5, d=1)])
        evs, _ = self.run_files([f], [0.0])
        v = self.valid(evs)
        self.assertEqual(len(v), 1)
        self.assertTrue(v[0]['end_known']); self.assertFalse(v[0]['need_check'])

    def test67a_stop_longer_than_max_static(self):
        """#67 停得比 max_static（測試用 15 秒）久，之後開走 → 只有一筆，包含整段停車，車尾已確認"""
        d, f, exp = self.scene('a.avi', 150, trains=[dict(t0=60, d=1, stops=[(61.0, 30)])])
        prof = core.Profile(); prof.max_static = 15
        evs, _ = self.run_files([f], [0.0], prof)
        v = self.valid(evs)
        self.assertEqual(len(v), 1, [(e['start'], e['end'], e['reason']) for e in v])
        self.assertLess(abs(v[0]['start'] - exp[0][0]), 0.3)
        self.assertLess(abs(v[0]['end'] - exp[0][1]), 0.4)
        self.assertTrue(v[0]['end_known'])
        self.assertTrue(v[0]['need_check'])
        self.assertIn('長時間', v[0]['reason'])

    # test67b／67c（攝影機被碰歪後重建背景繼續判讀）已由 v1.0.11 的 TestRound4.test71a／71b（停止判讀）取代（#71）

    def test68_time_jump_in_middle_file(self):
        """#68 事件跨 A、B、C，只有中間的 B 畫面時間跳動 → 也要標需確認"""
        try:
            from trainwatch import gui
        except Exception as e:
            self.skipTest(str(e))
        ev = dict(valid=True, start_file='A', end_file='C', reason='', need_check=False)
        gui.mark_time_jumps([ev], ['A', 'B', 'C'], [dict(msg=''), dict(msg='影片中的畫面時間有跳動'), dict(msg='')])
        self.assertTrue(ev['need_check'])

    def test69_time_goes_back(self):
        """#69 實際讀完 A 是 0～20 秒，B 的開始時間卻是 15 秒（倒退／重疊）→ 停止判讀並說明"""
        d = _mkdtemp()
        a, b = os.path.join(d, 'a.avi'), os.path.join(d, 'b.avi')
        selftest.make_scene(a, 20); selftest.make_scene(b, 20)
        with self.assertRaises(core.TimeOrderError) as cm:
            core.process([a, b], [0.0, 15.0], core.Profile())
        self.assertIn('b.avi', str(cm.exception))

    def test69b_continuous_files_ok(self):
        """#69 正常連續（誤差 1 秒內）的影片照常判讀，跨檔案的列車仍是一筆"""
        d = _mkdtemp()
        a, b = os.path.join(d, 'a.avi'), os.path.join(d, 'b.avi')
        selftest.make_scene(a, 20, trains=[dict(t0=18, d=1)])
        selftest.make_scene(b, 20, trains=[dict(t0=-2, d=1)])
        evs, _ = self.run_files([a, b], [0.0, 19.6])
        v = self.valid(evs)
        self.assertEqual(len(v), 1, [(e['start'], e['end']) for e in v])
        self.assertTrue(v[0]['end_known'])

    def test70_shot_names_when_unknown(self):
        """#70 車尾未確認時最後一張叫「3最後畫面_車尾未確認」；影片一開始就有車時第一張叫「1第一畫面_車頭未確認」"""
        d, f, exp = self.scene('a.avi', 20, trains=[dict(t0=-1, d=1), dict(t0=16, d=1)])
        evs, prof = self.run_files([f], [0.0])
        out = os.path.join(d, 'out')
        core.save_frames_and_clips(evs, prof, out, files=[f], bases=[0.0])
        v = self.valid(evs)
        self.assertEqual(len(v), 2, [(e['start'], e['end']) for e in v])
        self.assertTrue(any('1第一畫面_車頭未確認' in s for s in v[0]['shots']), v[0]['shots'])
        self.assertTrue(any('3最後畫面_車尾未確認' in s for s in v[1]['shots']), v[1]['shots'])


class TestRound4(unittest.TestCase):
    """v1.0.11：攝影機位置改變（修正清單 #71）。v1.0.13 起改成繼續判讀，之後的列車標需確認（#88）"""
    scene = TestRound2.scene
    valid = TestRound2.valid

    def run_info(self, files, bases, prof=None):
        info = {}
        evs = core.number_events(core.process(files, bases, prof or core.Profile(), info=info))
        return evs, info

    def test71a_idle_camera_moved_stops(self):
        """#71／#88 沒有列車時攝影機被碰歪 → 偵測到、回報影片與時間；**繼續判讀**，之後的列車判讀得到、標需確認並寫原因"""
        d, f, exp = self.scene('a.avi', 240, trains=[dict(t0=200, d=1)], camera_shift=(120, 14, 9, 1.5), texture=3000)
        evs, info = self.run_info([f], [0.0])
        self.assertIn('camera', info, info)
        self.assertEqual(info['camera']['file'], f)
        self.assertLess(abs(info['camera']['t'] - 120), 1.0, info)
        after = [e for e in self.valid(evs) if e['start'] > 121]
        self.assertEqual(len(after), 1, [(e['start'], e['valid'], e['reason'][:60]) for e in evs])
        self.assertLess(abs(after[0]['start'] - 200), 0.5)
        self.assertTrue(after[0]['need_check'])
        self.assertIn('攝影機位置', after[0]['reason'])
        self.assertEqual([e for e in self.valid(evs) if 115 < e['start'] < 135], [], '攝影機移動那一刻不可以多出一筆列車')

    def test71b_camera_moved_during_train_stops(self):
        """#71／#88 列車通過途中攝影機被碰歪 → 那一筆保留、車尾未確認、需確認；繼續判讀，之後的列車判讀得到、標需確認"""
        d, f, exp = self.scene('a.avi', 260, trains=[dict(t0=150, d=1, speed=40), dict(t0=220, d=-1)],
                               camera_shift=(151, 14, 9, 1.5), texture=3000)
        evs, info = self.run_info([f], [0.0])
        v = self.valid(evs)
        info2 = [(round(e['start'], 1), round(e['end'], 1), e['valid'], e['reason'][:50]) for e in evs]
        self.assertIn('camera', info, info2)
        self.assertEqual(len(v), 2, info2)
        self.assertLess(abs(v[0]['start'] - 150), 0.5)
        self.assertFalse(v[0]['end_known']); self.assertTrue(v[0]['need_check'])
        self.assertIn('攝影機位置', v[0]['reason'])
        self.assertLess(abs(v[1]['start'] - 220), 0.5, info2)
        self.assertTrue(v[1]['need_check'])
        self.assertIn('攝影機位置', v[1]['reason'])

    def test71c_no_false_camera_stop(self):
        """#71 曝光變化、車燈、紅外線切換、陽光、一般列車 → 不可以誤判成攝影機移動"""
        cases = [dict(seconds=60, trains=[dict(t0=20, d=1)], exposure=[(30, 45, 0.6)]),
                 dict(seconds=40, blobs=[(10, 14, 350, 190, 60), (20, 23, 300, 150, 90)]),
                 dict(seconds=60, ir_at=20, trains=[dict(t0=35, d=1)]),
                 dict(seconds=60, local_light=[(10, 50, 1.35)], trains=[dict(t0=25, d=-1)]),
                 dict(seconds=60, trains=[dict(t0=10, d=1, speed=40, length=800, y=(166, 188)),
                                          dict(t0=20, d=-1, speed=40, length=800, y=(192, 214))])]
        for kw in cases:
            sec = kw.pop('seconds')
            d, f, exp = self.scene('a.avi', sec, **kw)
            evs, info = self.run_info([f], [0.0])
            self.assertNotIn('camera', info, kw)


class TestRound5(unittest.TestCase):
    """v1.0.12：攝影機位置基準和列車背景分開（修正清單 #75～#77）"""
    scene = TestRound2.scene
    valid = TestRound2.valid
    run_info = TestRound4.run_info

    def test75a_move_at_30s_inside_initial_window(self):
        """#75 第 30 秒攝影機永久移動（還在開頭 90 秒的背景取樣區間內）→ 一定要偵測到；之後的列車判讀得到、標需確認（#88）"""
        d, f, exp = self.scene('a.avi', 150, trains=[dict(t0=120, d=1)], camera_shift=(30, 14, 9, 1.5), texture=3000)
        evs, info = self.run_info([f], [0.0])
        self.assertIn('camera', info, [(e['start'], e['end'], e['valid'], e['reason'][:40]) for e in evs])
        self.assertLess(abs(info['camera']['t'] - 30), 1.5, info)
        v = self.valid(evs)
        self.assertEqual(len(v), 1, [(e['start'], e['end'], e['valid'], e['reason'][:60]) for e in evs])
        self.assertLess(abs(v[0]['start'] - 120), 0.5)
        self.assertTrue(v[0]['need_check'])

    def test75b_move_at_3s(self):
        """#75 第 3 秒攝影機永久移動 → 偵測到；移動以前那幾秒不可以變成一列「列車」"""
        d, f, exp = self.scene('a.avi', 60, camera_shift=(3, 14, 9, 1.5), texture=3000)
        evs, info = self.run_info([f], [0.0])
        self.assertIn('camera', info, [(e['start'], e['end'], e['valid'], e['reason'][:40]) for e in evs])
        self.assertEqual(self.valid(evs), [], [(e['start'], e['end'], e['reason'][:60]) for e in evs])

    def test75c_train_at_start_no_move(self):
        """#75 影片一開頭就有列車通過、攝影機沒動 → 不可以誤判成攝影機移動"""
        for t0 in (-1, 0.5, 3):
            d, f, exp = self.scene('a.avi', 40, trains=[dict(t0=t0, d=1)], texture=3000)
            evs, info = self.run_info([f], [0.0])
            self.assertNotIn('camera', info, t0)

    def gap_case(self, b_kw):
        d = _mkdtemp()
        a, b = os.path.join(d, 'a.avi'), os.path.join(d, 'b.avi')
        selftest.make_scene(a, 30, texture=3000)
        selftest.make_scene(b, 40, trains=[dict(t0=20, d=1)], texture=3000, **b_kw)
        return self.run_info([a, b], [0.0, 60.0]) + (b,)

    def test76d_gap_same_place_brightness(self):
        """#76 錄影中斷前後攝影機沒動、只是整體亮度不同 → 照常繼續，中斷後的列車判讀得到"""
        evs, info, b = self.gap_case(dict(exposure=[(0, 40, 0.7)]))
        self.assertNotIn('camera', info, info)
        self.assertEqual(len(self.valid(evs)), 1)

    def test76e_gap_camera_moved(self):
        """#76 錄影中斷期間攝影機被移動 → 回報影片、時間，說明中斷前後位置不一致；繼續判讀，之後的列車標需確認（#88）"""
        evs, info, b = self.gap_case(dict(camera_shift=(0, 14, 9, 1.5)))
        self.assertIn('camera', info, info)
        self.assertEqual(info['camera']['file'], b)
        self.assertEqual(info['camera'].get('kind'), 'gap')
        v = self.valid(evs)
        self.assertEqual(len(v), 1)
        self.assertTrue(v[0]['need_check'])
        self.assertIn('錄影中斷後', v[0]['reason'])

    def test76f_gap_day_to_ir(self):
        """#76 中斷前白天、中斷後紅外線，攝影機沒動 → 能確認位置一致就照常；確認不了就說明「無法確認」、之後的列車標需確認，
        不可以默默當成位置一定沒變。列車都要判讀得到（#88）"""
        evs, info, b = self.gap_case(dict(ir_at=0))
        self.assertEqual(len(self.valid(evs)), 1)
        if 'camera' in info:
            self.assertEqual(info['camera'].get('kind'), 'gap')
            self.assertIn(info['camera'].get('result'), ('uncertain', 'moved'))
            self.assertTrue(self.valid(evs)[0]['need_check'])


class TestRound6(unittest.TestCase):
    """v1.0.12：交會的各種情形（使用者 2026-10-02 確認的規則，#79）"""
    scene = TestRound2.scene
    valid = TestRound2.valid
    run_info = TestRound4.run_info

    def test79a_crossing_a_longer(self):
        """A 比 B 長（A 110～140 秒，B 在 A 期間 118～128 秒）→ 一筆＝A 的進入、離開"""
        d, f, _ = self.scene('a.avi', 160, trains=[dict(t0=110, d=1, speed=40, length=1200, y=(166, 188)),
                                                     dict(t0=118, d=-1, speed=40, length=400, y=(192, 214))])
        evs, info = self.run_info([f], [0.0])
        v = self.valid(evs)
        self.assertEqual(len(v), 1, [(e['start'], e['end']) for e in v])
        self.assertLess(abs(v[0]['start'] - 110), 0.3); self.assertLess(abs(v[0]['end'] - 140), 0.4)
        self.assertTrue(v[0]['end_known'])

    def test79b_crossing_b_leaves_last(self):
        """A 先離開、B 還沒（A 110～130 秒，B 115～145 秒）→ 一筆＝A 的進入、B 的車尾離開"""
        d, f, _ = self.scene('a.avi', 160, trains=[dict(t0=110, d=1, speed=40, length=800, y=(166, 188)),
                                                     dict(t0=115, d=-1, speed=40, length=1200, y=(192, 214))])
        evs, info = self.run_info([f], [0.0])
        v = self.valid(evs)
        self.assertEqual(len(v), 1, [(e['start'], e['end']) for e in v])
        self.assertLess(abs(v[0]['start'] - 110), 0.3); self.assertLess(abs(v[0]['end'] - 145), 0.4)

    def test79c_extreme_b_stops(self):
        """最極端：A 進入一半時 B 來了、B 停在參考線上；A 開走後 B 還停著，最後 B 才開走
        → 一筆＝A 的進入～B 的車尾離開，含整段停車，車尾已確認"""
        d, f, _ = self.scene('a.avi', 150, trains=[dict(t0=60, d=1, speed=40, length=800, y=(166, 188)),
                                                     dict(t0=70, d=-1, speed=40, length=300, y=(192, 214), stops=[(72.0, 40)])])
        evs, info = self.run_info([f], [0.0])
        v = self.valid(evs)
        self.assertEqual(len(v), 1, [(e['start'], e['end'], e['reason'][:40]) for e in v])
        self.assertLess(abs(v[0]['start'] - 60), 0.3)
        self.assertLess(abs(v[0]['end'] - (70 + 300 / 40 + 40)), 0.5)
        self.assertTrue(v[0]['end_known'])
        self.assertIn('停止', v[0]['reason'])


class TestRound7(unittest.TestCase):
    """#80：錄影中斷停止的說法、位置基準重新建立的紀錄要寫出來"""

    def _gui(self):
        try:
            from trainwatch import gui
        except Exception as e:
            self.skipTest(str(e))
        return gui

    def test80a_event_text_kinds(self):
        """#80／#88 攝影機位置有問題的說明：各種情形說法不同，都寫影片、時間，並說明之後的列車標需確認"""
        gui = self._gui()
        t0 = dt.datetime(2030, 1, 1, 8, 0, 0).timestamp()
        ks = [dict(t=t0, file='x/a.mkv', pos=12.0, kind='moved', result='moved', disp=8.0, now=t0, warm_until=t0 + 10),
              dict(t=t0, file='x/b.mkv', pos=0.0, kind='gap', result='moved', prev_end=t0 - 60, disp=9.0),
              dict(t=t0, file='x/b.mkv', pos=0.0, kind='gap', result='uncertain', prev_end=t0 - 60),
              dict(t=t0, file='x/a.mkv', pos=30.0, kind='drift', result='moved', disp=6.5),
              dict(t=t0, file='x/a.mkv', pos=30.0, kind='uncertain', result='uncertain'),
              dict(t=t0, file='x/a.mkv', pos=0.0, kind='profile', result='moved', disp=12.0)]
        txt = [gui.camera_event_text(k) for k in ks]
        self.assertEqual(len(set(txt)), len(txt))
        for x, k in zip(txt, ks):
            self.assertIn(os.path.basename(k['file']), x)
            self.assertIn('08:00:00', x)
            self.assertIn('標需人工確認', x)
        self.assertIn('07:59:00', txt[1])          # 中斷前最後畫面時間
        self.assertIn('補判讀', txt[0])
        self.assertIn('8.0', txt[0])
        q = gui.quality_summary([], timing=[dict(warn=txt[0])])
        self.assertIn('已繼續判讀', q)
        self.assertIn('重畫參考線', q)
        self.assertIn('逐筆確認', q)

    def test80b_notes_shown(self):
        gui = self._gui()
        files = ['x/a.mkv', 'x/b.mkv']
        timing = [dict(offset=0), dict(offset=600)]
        info = dict(notes=[dict(file='x/b.mkv', t=700, text='08:11:40：切換為紅外線（黑白）（b.mkv）')])
        gui.apply_camera_info(info, files, timing)
        self.assertNotIn('cam_note', timing[0])
        self.assertIn('紅外線', timing[1]['cam_note'])
        self.assertNotIn('warn', timing[1])
        q = gui.quality_summary([], timing=timing)
        self.assertIn('位置基準重新建立', q)
        self.assertIn('08:11:40', q)
        d = _mkdtemp()
        p = os.path.join(d, 'r.xlsx')
        export.write_excel(p, [], core.Profile(), files, timing)
        import openpyxl
        ws = openpyxl.load_workbook(p)['設定與影片']
        cells = [str(c.value) for row in ws.iter_rows() for c in row if c.value]
        self.assertTrue(any('08:11:40' in c for c in cells), cells)
        info = dict(camera_events=[dict(t=1e9, file='x/a.mkv', pos=3.0, kind='moved', result='moved', now=1e9)])
        gui.apply_camera_info(info, files, timing)
        self.assertIn('之後的列車', timing[0]['warn'])
        export.write_excel(p, [], core.Profile(), files, timing)
        ws = openpyxl.load_workbook(p)['設定與影片']
        cells = [str(c.value) for row in ws.iter_rows() for c in row if c.value]
        self.assertTrue(any('判讀途中的狀況' in c for c in cells), cells)


class TestRound8(unittest.TestCase):
    """v1.0.13：監測站的攝影機位置基準（#82）、慢慢偏移（#83）、無法確認時停止（#84）、錄影中斷沒有基準（#85）、
    可接受誤差 1.5 秒（#86）"""
    scene = TestRound2.scene
    valid = TestRound2.valid
    run_info = TestRound4.run_info

    def anchored(self, f, at=1.0):
        """用 f 第 at 秒的畫面建立位置基準（等於在這個畫面畫好參考線、儲存監測站）"""
        import cv2
        c = cv2.VideoCapture(f)
        c.set(cv2.CAP_PROP_POS_MSEC, at * 1000)
        ok, fr = c.read()
        c.release()
        p = core.Profile()
        p.camera_anchors = [core.make_anchor(fr)]
        return p

    def test82a_profile_anchor_same_place(self):
        """#82 監測站在位置 A 存基準；下一批影片也是位置 A → 照常判讀"""
        d, a, _ = self.scene('a.avi', 20, texture=3000)
        d, b, _ = self.scene('b.avi', 40, trains=[dict(t0=20, d=1)], texture=3000, seed=2)
        evs, info = self.run_info([b], [0.0], self.anchored(a))
        self.assertNotIn('camera', info, info)
        self.assertEqual(len(self.valid(evs)), 1)

    def test82b_moved_before_batch(self):
        """#82 監測站在位置 A 存基準；下一批影片從第一格起就移動 14×9＋轉 1.5 度 → 判讀開始時就發現（kind=profile）；
        使用者選「仍要判讀」時（#88）照常判讀，全部列車標需確認並寫原因"""
        d, a, _ = self.scene('a.avi', 20, texture=3000)
        d, b, _ = self.scene('b.avi', 40, trains=[dict(t0=20, d=1)], texture=3000, camera_shift=(0, 14, 9, 1.5))
        p = self.anchored(a)
        self.assertEqual([x['verdict'] for x in core.check_start([b], p)], ['moved'])     # 畫面會先問使用者
        evs, info = self.run_info([b], [0.0], p)
        self.assertIn('camera', info, info)
        self.assertEqual(info['camera'].get('kind'), 'profile')
        self.assertEqual(info['camera'].get('result'), 'moved')
        self.assertLess(info['camera']['t'], 1.0)
        v = self.valid(evs)
        self.assertEqual(len(v), 1)
        self.assertTrue(v[0]['need_check'])
        self.assertIn('監測站儲存時不同', v[0]['reason'])

    def test82c_anchor_brightness(self):
        """#82 隔天亮度差很多、攝影機沒動 → 位置相同；確認不了也只能是「無法確認」，不可以判成移動"""
        d, a, _ = self.scene('a.avi', 20, texture=3000)
        d, b, _ = self.scene('b.avi', 30, texture=3000, exposure=[(0, 30, 0.6)])
        evs, info = self.run_info([b], [0.0], self.anchored(a))
        if 'camera' in info:
            self.assertEqual(info['camera'].get('result'), 'uncertain', info)

    def test82d_check_start(self):
        """#82 判讀前檢查：沒有基準（舊版監測站）→ 要使用者確認；基準位置不同 → moved；
        批次中途畫面變成紅外線 → 紅外線那段也要檢查"""
        d, a, _ = self.scene('a.avi', 10, texture=3000)
        d, m, _ = self.scene('m.avi', 10, texture=3000, camera_shift=(0, 14, 9, 1.5))
        d, n, _ = self.scene('n.avi', 10, texture=3000, ir_at=0)
        r = core.check_start([a], core.Profile())
        self.assertEqual([x['verdict'] for x in r], ['none'])
        p = self.anchored(a)
        self.assertEqual([x['verdict'] for x in core.check_start([a], p)], ['same'])
        self.assertEqual([x['verdict'] for x in core.check_start([m], p)], ['moved'])
        r = core.check_start([a, n], p)
        self.assertEqual([x['file'] for x in r], [a, n])
        self.assertEqual(r[0]['verdict'], 'same')
        self.assertIn(r[1]['verdict'], ('same', 'uncertain'))

    def test83a_slow_drift(self):
        """#83 攝影機從第 5 秒起每 2 秒偏 0.6 像素（每次都小於基準更新門檻 1 像素）→ 累積超過 6 像素後要偵測到"""
        d, f, _ = self.scene('a.avi', 120, texture=3000, cam_path=[(0, 0, 0), (5, 0, 0), (95, 27, 0)])
        evs, info = self.run_info([f], [0.0])
        self.assertIn('camera', info, info)
        self.assertEqual(info['camera'].get('kind'), 'drift', info)
        self.assertLess(info['camera']['t'], 100, info)

    def test83b_jitter_no_false_alarm(self):
        """#83 攝影機每 2 秒隨機晃 ±0.5 像素、長期平均在原位 → 不可以誤判"""
        rng = np.random.default_rng(5)
        path = [(t, float(rng.uniform(-0.5, 0.5)), float(rng.uniform(-0.5, 0.5))) for t in range(0, 151, 2)]
        d, f, _ = self.scene('a.avi', 150, texture=3000, cam_path=path)
        evs, info = self.run_info([f], [0.0])
        self.assertNotIn('camera', info, info)

    def test83c_drift_and_back(self):
        """#83 先慢慢偏 4 像素、再慢慢回到原位 → 不可以因為「總共移動了 8 像素」就誤判"""
        d, f, _ = self.scene('a.avi', 150, texture=3000, cam_path=[(0, 0, 0), (40, 4, 0), (80, 0, 0)])
        evs, info = self.run_info([f], [0.0])
        self.assertNotIn('camera', info, info)

    def test84a_rotation(self):
        """#84 純轉動 4 度（平移幾乎是 0）→ 軌道範圍兩端偏移超過 6 像素，要偵測到"""
        d, f, _ = self.scene('a.avi', 90, texture=3000, camera_shift=(20, 0, 0, 4.0))
        evs, info = self.run_info([f], [0.0])
        self.assertIn('camera', info, info)

    def test84b_zoom(self):
        """#84 變焦 5%（平移是 0）→ 軌道範圍兩端偏移超過 6 像素，要偵測到"""
        d, f, _ = self.scene('a.avi', 90, texture=3000, cam_path=[(0, 0, 0), (20, 0, 0, 0, 1.0), (20.1, 0, 0, 0, 1.05)])
        evs, info = self.run_info([f], [0.0])
        self.assertIn('camera', info, info)

    def test84c_ir_switch_no_stop(self):
        """#84 攝影機沒動、第 20 秒切換紅外線 → 不可以因為切換就當成位置有問題"""
        d, f, _ = self.scene('a.avi', 150, texture=3000, ir_at=20, trains=[dict(t0=120, d=1)])
        evs, info = self.run_info([f], [0.0])
        self.assertNotIn('camera', info, info)
        self.assertEqual(len(self.valid(evs)), 1)

    def test84d_uncertain_then_recovered(self):
        """#88 畫面有一段時間比不起來（例如傍晚轉換、大雨），之後又和監測站基準確認位置相同 →
        只有比不起來那段時間的列車標需確認，之後的列車照常（不是從此全部標需確認）"""
        import cv2
        d = _mkdtemp()
        src, f = os.path.join(d, 's.avi'), os.path.join(d, 'f.avi')
        selftest.make_scene(src, 150, trains=[dict(t0=120, d=1)], texture=3000)
        c = cv2.VideoCapture(src)
        w = cv2.VideoWriter(f, cv2.VideoWriter_fourcc(*'MJPG'), selftest.FPS, (selftest.W, selftest.H))
        rng = np.random.default_rng(1)
        i = 0
        while True:
            ok, fr = c.read()
            if not ok:
                break
            if 20 * selftest.FPS <= i < 60 * selftest.FPS:       # 這 40 秒畫面完全比不起來
                fr = rng.integers(0, 255, fr.shape, np.uint8)
            w.write(fr); i += 1
        w.release(); c.release()
        old = core.GEO_STALE
        core.GEO_STALE = 10.0
        try:
            evs, info = self.run_info([f], [0.0], self.anchored(src))
        finally:
            core.GEO_STALE = old
        ces = info.get('camera_events', [])
        self.assertTrue(any(x['kind'] == 'uncertain' and x.get('until') for x in ces), ces)
        v = [e for e in self.valid(evs) if abs(e['start'] - 120) < 1]
        self.assertEqual(len(v), 1, [(e['start'], e['valid'], e['reason'][:50]) for e in evs])
        self.assertNotIn('攝影機', v[0]['reason'])

    def test85_gap_without_reference(self):
        """#85 還沒有位置基準就遇到錄影中斷 → 不可以回答「相同」"""
        d, f, _ = self.scene('a.avi', 5, texture=3000)
        g = core.CameraGuard(core.Profile())
        self.assertEqual(g.gap_check(f), 'uncertain')

    def test87_error_keeps_results(self):
        """#87／#88 第二支影片判讀途中出錯 → 第一支的列車照常回傳；跳過第二支剩下的部分、**繼續判讀第三支**；
        回報在哪一支、哪個時間出錯"""
        d = _mkdtemp()
        a, b, c = (os.path.join(d, x) for x in ('a.avi', 'b.avi', 'c.avi'))
        for x in (a, b, c):
            selftest.make_scene(x, 30, trains=[dict(t0=10, d=1)])
        orig = core.Detector.feed

        def boom(self, t_abs, *a_, **k):
            if 35 <= t_abs < 60:
                raise ValueError('模擬解碼錯誤')
            return orig(self, t_abs, *a_, **k)
        core.Detector.feed = boom
        try:
            info = {}
            evs = core.process([a, b, c], [0.0, 30.0, 60.0], core.Profile(), info=info)
        finally:
            core.Detector.feed = orig
        v = [e for e in evs if e['valid']]
        self.assertEqual([round(e['start']) for e in v], [10, 70], [(e['start'], e['reason'][:50]) for e in evs])
        self.assertEqual(info['error']['file'], b)
        self.assertIn('模擬解碼錯誤', info['error']['msg'])
        try:
            from trainwatch import gui
        except Exception:
            return
        timing = [dict(offset=0), dict(offset=30), dict(offset=60)]
        gui.apply_camera_info(info, [a, b, c], timing)
        self.assertIn('讀取出錯', timing[1]['warn'])
        self.assertIn('從下一支繼續', timing[1]['warn'])

    def test86_time_tolerance(self):
        """#86 使用者可接受的誤差：車頭進入、車尾離開參考線 1.5 秒以內 → 自動修正超過 1.5 秒要標需確認"""
        self.assertEqual(core.REFINE_CHECK, 1.5)
        import inspect
        self.assertNotIn('可接受誤差是 5 秒', inspect.getsource(core))


class TestRound2Gui(unittest.TestCase):
    """v1.0.9：介面這邊的檢查函式（不開視窗）"""
    def setUp(self):
        try:
            from trainwatch import gui
        except Exception as e:
            self.skipTest(str(e))
        self.gui = gui

    def test48_final_order(self):
        """#48 正式校正後順序和清單不同 → 依時間排序；重疊超過 2 秒 → 列出來（不能判讀）"""
        d, f, _ = TestRound2.scene(self, 'a.avi', 10)
        files = [f + '1', f + '2', f + '3']
        for x in files:
            os.link(f, x) if not os.path.exists(x) else None
        tm = [dict(offset=1000.0), dict(offset=0.0), dict(offset=500.0)]
        fs, t2, note, ov = self.gui.final_order(files, tm)
        self.assertEqual([t['offset'] for t in t2], [0.0, 500.0, 1000.0])
        self.assertEqual(fs, [files[1], files[2], files[0]])
        self.assertTrue(note)
        self.assertEqual(ov, [])
        fs, t2, note, ov = self.gui.final_order(files, [dict(offset=0.0), dict(offset=5.0), dict(offset=500.0)])
        self.assertEqual(len(ov), 1)            # 10 秒的影片，第二支 5 秒就開始 → 重疊
        self.assertEqual(note, '')

    def test52_time_jump_marks_events(self):
        """#52 畫面時間跳動的影片：裡面的列車標需確認，備註寫原因"""
        evs = [dict(valid=True, start_file='a', end_file='a', reason='', need_check=False),
               dict(valid=True, start_file='b', end_file='b', reason='', need_check=False)]
        self.gui.mark_time_jumps(evs, ['a', 'b'], [dict(msg='影片中的畫面時間有跳動（9 個取樣對不上），請確認'), dict(msg='')])
        self.assertTrue(evs[0]['need_check']); self.assertIn('畫面時間有跳動', evs[0]['reason'])
        self.assertFalse(evs[1]['need_check'])

    def test61_quality_summary(self):
        """#61 品質摘要：需確認的原因分類計數、車尾未確認另外提醒"""
        evs = [dict(valid=True, need_check=True, check_items=['方向不確定']),
               dict(valid=True, need_check=True, end_known=False, check_items=['方向不確定', '影片結束時列車仍在參考線上（車尾未確認）']),
               dict(valid=True, need_check=False), dict(valid=False)]
        t = self.gui.quality_summary(evs)
        self.assertIn('列車 3 筆、已排除 1 筆', t)
        self.assertIn('方向不確定：2 筆', t)
        self.assertIn('影片結束時列車仍在參考線上（車尾未確認）：1 筆', t)
        self.assertIn('1筆「車尾未確認」', t)


class TestExport(unittest.TestCase):
    def test_noise_txt_midnight(self):
        """#44 跨午夜的列車：一行，日期用車頭進入那天（使用者實測噪音分析程式可以接受）"""
        b = dt.datetime(2023, 9, 18, 23, 59, 58).timestamp()
        self.assertEqual(export.noise_lines([dict(start=b + 0.2, end=b + 4.6, valid=True)]),
                         ['2023/09/18,23:59:58,00:00:03,10,FR,,A'])

    def test_noise_txt(self):
        d = _mkdtemp()
        b = dt.datetime(2023, 9, 18, 11, 1, 26).timestamp()
        evs = [dict(start=b + 0.7, end=b + 2.2, valid=True),
               dict(start=b + 300, end=b + 305, valid=False),                 # 非列車不可以出現
               dict(start=b + 510.0, end=b + 516.0, valid=True),
               dict(start=b - 100, end=b - 90, valid=True)]                   # 順序要依時間排
        p = os.path.join(d, 'x.txt')
        self.assertEqual(export.write_noise_txt(p, evs), 3)
        with open(p, 'rb') as fh:
            raw = fh.read()
        self.assertEqual(raw, b'2023/09/18,10:59:46,10:59:56,10,FR,,A\r\n'
                              b'2023/09/18,11:01:26,11:01:29,10,FR,,A\r\n'
                              b'2023/09/18,11:09:56,11:10:02,10,FR,,A\r\n')

    def test_roundtrip(self):
        d = _mkdtemp()
        base = dt.datetime(2026, 1, 2, 10, 0, 0).timestamp()
        evs = [dict(no=1, start=base + 10, end=base + 25, start_file='a.mkv', end_file='a.mkv', start_pos=10, end_pos=25,
                    direction='往右', valid=True, need_check=True, reason='方向不確定', shots=['001_1車頭.jpg'], clip='001.mp4'),
               dict(no=1, start=base + 40, end=base + 40.2, start_file='a.mkv', end_file='a.mkv', start_pos=40, end_pos=40.2,
                    direction='無法判定', valid=False, reason='時間太短')]
        prof = core.Profile()
        timing = [dict(offset=base, source='畫面字幕', support=1.0, msg='')]
        export.save_json(os.path.join(d, 'results.json'), evs, prof, ['a.mkv'], timing)
        e2, p2, f2, t2 = export.load_json(os.path.join(d, 'results.json'))
        self.assertEqual(len(e2), 2)
        self.assertEqual(e2[0]['direction'], '往右')
        export.write_excel(os.path.join(d, 'x.xlsx'), e2, p2, f2, t2)
        export.write_csv(os.path.join(d, 'x.csv'), e2)
        from openpyxl import load_workbook
        wb = load_workbook(os.path.join(d, 'x.xlsx'))
        self.assertEqual(wb.sheetnames, ['列車通過紀錄', '已排除', '設定與影片'])
        ws = wb['列車通過紀錄']
        self.assertEqual(ws.cell(2, 3).value, '10:00:10.0')
        self.assertEqual(ws.cell(2, 9).value, '是')
        self.assertEqual(wb['已排除'].max_row, 2)
        with open(os.path.join(d, 'x.csv'), encoding='utf-8-sig') as fh:
            self.assertEqual(len(fh.read().strip().splitlines()), 2)


class TestGuiHelpers(unittest.TestCase):
    def test_parse(self):
        try:
            from trainwatch import gui
        except Exception as e:  # 沒有 tkinter 的環境
            self.skipTest(str(e))
        self.assertAlmostEqual(gui.parse_hms('08:15:30.5'), 29730.5)
        self.assertAlmostEqual(gui.parse_hms('081530'), 29730)
        self.assertEqual(gui.parse_datetime('2026/01/02 08:00:00'), dt.datetime(2026, 1, 2, 8, 0, 0))
        with self.assertRaises(ValueError):
            gui.parse_hms('25:00:00')

    def test_result_location(self):
        """結果統一放進「判讀結果」資料夾；預設在影片旁邊；選的資料夾本身叫「判讀結果」時不再多包一層；預設位置存得起來"""
        try:
            from trainwatch import gui
        except Exception as e:
            self.skipTest(str(e))
        import tempfile
        v = os.path.join('D:', os.sep, '錄影', 'a.mkv')
        self.assertEqual(gui.result_root(None, v), os.path.join('D:', os.sep, '錄影', '判讀結果'))
        self.assertEqual(gui.result_root(os.path.join('E:', os.sep, '報告'), v), os.path.join('E:', os.sep, '報告', '判讀結果'))
        self.assertEqual(gui.result_root(os.path.join('E:', os.sep, '判讀結果'), v), os.path.join('E:', os.sep, '判讀結果'))
        with tempfile.TemporaryDirectory() as d:
            old = os.environ.get('TRAINWATCH_DATA')
            os.environ['TRAINWATCH_DATA'] = d
            try:
                self.assertEqual(gui.load_settings(), {})
                self.assertTrue(gui.save_settings({'result_location': os.path.join(d, '報告')}))
                self.assertEqual(gui.load_settings()['result_location'], os.path.join(d, '報告'))
                self.assertTrue(os.path.exists(os.path.join(d, '程式設定.json')))
            finally:
                if old is None:
                    os.environ.pop('TRAINWATCH_DATA')
                else:
                    os.environ['TRAINWATCH_DATA'] = old

    def test_import_old_settings(self):
        """換新版：從舊版程式資料夾（或直接選「監測站設定」資料夾）找到設定並複製；同名的不覆蓋"""
        try:
            from trainwatch import gui
        except Exception as e:
            self.skipTest(str(e))
        import tempfile, json as _j
        with tempfile.TemporaryDirectory() as d:
            old = os.path.join(d, '列車通過判讀_v1.0.4')
            os.makedirs(os.path.join(old, '監測站設定'))
            for n in ('甲站', '乙站'):
                with open(os.path.join(old, '監測站設定', n + '.json'), 'w', encoding='utf-8') as fh:
                    fh.write('{"name": "%s"}' % n)
            with open(os.path.join(old, '監測站設定', '_上次使用.txt'), 'w', encoding='utf-8') as fh:
                fh.write('乙站')
            with open(os.path.join(old, '程式設定.json'), 'w', encoding='utf-8') as fh:
                _j.dump({'result_location': 'E:\\報告'}, fh)
            for pick in (old, os.path.join(old, '監測站設定')):
                prof, st = gui.find_old_settings(pick)
                self.assertEqual(prof, os.path.join(old, '監測站設定'))
                self.assertEqual(st, os.path.join(old, '程式設定.json'))
            self.assertEqual(gui.find_old_settings(d), (None, None))
            new = os.path.join(d, '文件', '列車通過判讀設定')
            os.makedirs(os.path.join(new, '監測站設定'))
            with open(os.path.join(new, '監測站設定', '甲站.json'), 'w', encoding='utf-8') as fh:
                fh.write('{"name": "新的甲站"}')
            copied, skipped, st_done = gui.import_settings(*gui.find_old_settings(old), new)
            self.assertEqual(copied, ['乙站']); self.assertEqual(skipped, ['甲站']); self.assertTrue(st_done)
            with open(os.path.join(new, '監測站設定', '甲站.json'), encoding='utf-8') as fh:
                self.assertIn('新的甲站', fh.read())          # 同名的沒有被舊版蓋掉
            self.assertTrue(os.path.exists(os.path.join(new, '程式設定.json')))
            self.assertEqual(gui.profile_names_in(os.path.join(new, '監測站設定')), ['乙站', '甲站'])


def _expected():
    root = os.environ.get('TW_SAMPLES')
    p = os.path.join(root, 'expected.json') if root else ''
    return p if p and os.path.exists(p) else None


@unittest.skipUnless(_expected(), '沒有設定 TW_SAMPLES，或樣本資料夾裡沒有 expected.json')
class TestSamples(unittest.TestCase):
    """樣本資料夾放影片＋expected.json（逐格截圖核對過的答案）。兩者都是真實監測資料，不放進 repository。
    expected.json 格式：[{"files": ["檔名開頭", ...], "events": [{"start": "時:分:秒.x", "end": "…", "direction": "往右"|null}]}]"""
    def test_samples(self):
        root = os.environ['TW_SAMPLES']
        import json
        with open(_expected(), encoding='utf-8') as fh:
            cases = json.load(fh)

        def sec(s):
            h, m, x = s.split(':')
            return int(h) * 3600 + int(m) * 60 + float(x)
        for case in cases:
            keys = case['files']
            exp = [(x['start'], x['end'], x['direction']) for x in case['events']]
            files = [glob.glob(os.path.join(root, k + '*'))[0] for k in keys]
            bases = []
            for f in files:
                r = osd.calibrate(f, osd.DEFAULT_RECT)
                self.assertTrue(r['ok'], (f, r))
                self.assertGreaterEqual(r['support'], 0.95)
                bases.append(r['offset'])
            info = {}
            evs = [e for e in core.number_events(core.process(files, bases, core.Profile(), info=info)) if e['valid']]
            self.assertNotIn('camera', info, (keys, info))          # 真實樣本不可以誤判成攝影機移動（#71）
            got = [(core.fmt_time(e['start']), core.fmt_time(e['end']), e['direction']) for e in evs]
            self.assertEqual(len(got), len(exp), (keys, got))
            for g, x in zip(got, exp):
                self.assertLess(abs(sec(g[0]) - sec(x[0])), 0.5, (keys, g, x))
                self.assertLess(abs(sec(g[1]) - sec(x[1])), 0.5, (keys, g, x))
                if x[2]:
                    self.assertEqual(g[2], x[2], (keys, g, x))


    def test_no_false_long_events(self):
        """樣本影片裡沒有停駛的列車，所以每一支單獨判讀時，不可以出現超過 30 秒的「列車」。
        （v1.0.4 開發時：初始背景改用整支影片的中位數，傍晚那支一開頭就被誤判成 54 秒的列車）"""
        root = os.environ['TW_SAMPLES']
        for f in sorted(glob.glob(os.path.join(root, '*.mkv'))):
            r = osd.calibrate(f, osd.DEFAULT_RECT)
            for e in core.process([f], [r['offset']], core.Profile()):
                if e['valid']:
                    self.assertLess(e['end'] - e['start'], 30, (f, core.fmt_time(e['start']), core.fmt_time(e['end'])))

    def test77_real_video_camera_moved(self):
        """#77 真實樣本（白天、夜間各一支）的前 120 秒，從第 60 秒起把整個畫面移動／轉動（模擬攝影機被碰歪）
        → 要偵測到並停止。v1.0.11 的偵測方式在真實畫面上偵測不到（假影片的圓點紋理才偵測得到）"""
        import cv2
        root = os.environ['TW_SAMPLES']
        allf = sorted(glob.glob(os.path.join(root, '*.mkv')))

        def sat(f):
            c = cv2.VideoCapture(f); c.set(cv2.CAP_PROP_POS_FRAMES, 30); ok, fr = c.read(); c.release()
            return cv2.cvtColor(fr, cv2.COLOR_BGR2HSV)[..., 1].mean()
        # 不在原始碼寫樣本的檔名（檔名就是真實時間）：白天＝彩度最高的一支，夜間＝排序最後的一支（紅外線）
        day = max(allf, key=sat)
        night = allf[-1]
        self.assertLess(sat(night), 5)
        for key, src in (('day', day), ('night', night)):
            for dx, dy, ang in ((14, 9, 1.5), (10, 0, 0)):
                d = _mkdtemp(); f = os.path.join(d, 'm.avi')
                c = cv2.VideoCapture(src)
                w = cv2.VideoWriter(f, cv2.VideoWriter_fourcc(*'MJPG'), 15.0, (640, 360))
                n = 0
                while n < 15 * 120:
                    ok, fr = c.read()
                    if not ok:
                        break
                    if n >= 15 * 60:
                        M = cv2.getRotationMatrix2D((320, 180), ang, 1.0); M[0, 2] += dx; M[1, 2] += dy
                        fr = cv2.warpAffine(fr, M, (640, 360), borderMode=cv2.BORDER_REFLECT)
                    w.write(fr); n += 1
                w.release(); c.release()
                info = {}
                core.process([f], [0.0], core.Profile(), info=info)
                self.assertIn('camera', info, (key, dx, dy, ang))
                self.assertLess(abs(info['camera']['t'] - 60), 2.0, (key, info))

    def test82_real_profile_anchor(self):
        """#82 真實樣本：用前一支影片第 1 秒存位置基準（等於當時畫好參考線、儲存監測站），
        下一支（白天、夜間各一組）的前 120 秒 → 位置相同、不可以有警告；同一段畫面整個移動 14×9＋轉 1.5 度 → 判讀開始時就發現
        （kind=profile），之後的列車都標需確認"""
        import cv2
        root = os.environ['TW_SAMPLES']
        allf = sorted(glob.glob(os.path.join(root, '*.mkv')))
        for a, b in ((allf[0], allf[1]), (allf[-2], allf[-1])):     # 不寫檔名（檔名就是真實時間）
            p = core.Profile()
            fr, _ = core.geo_frame_at(a)
            p.camera_anchors = [core.make_anchor(fr)]
            for move in (False, True):
                d = _mkdtemp(); f = os.path.join(d, 'm.avi')
                c = cv2.VideoCapture(b)
                w = cv2.VideoWriter(f, cv2.VideoWriter_fourcc(*'MJPG'), 15.0, (640, 360))
                for _ in range(15 * 120):
                    ok, x = c.read()
                    if not ok:
                        break
                    if move:
                        M = cv2.getRotationMatrix2D((320, 180), 1.5, 1.0); M[0, 2] += 14; M[1, 2] += 9
                        x = cv2.warpAffine(x, M, (640, 360), borderMode=cv2.BORDER_REFLECT)
                    w.write(x)
                w.release(); c.release()
                info = {}
                evs = core.process([f], [0.0], p, info=info)
                if move:
                    self.assertEqual(info.get('camera', {}).get('kind'), 'profile', info)
                    self.assertEqual(info['camera'].get('result'), 'moved', info)
                    self.assertTrue(all(e['need_check'] for e in evs if e['valid']))
                else:
                    self.assertNotIn('camera', info, info)

    def test_shots_and_clips_every_sample(self):
        """每支樣本影片（含每秒實際格數和檔頭不符的）：每一筆都要有 3 張截圖；列車要有短片，已排除的不產生短片（#60），
        而且車頭截圖上的畫面時間要和判讀的車頭時間一致（差 1 秒內）"""
        root = os.environ['TW_SAMPLES']
        import cv2
        T = osd.default_templates()
        for f in sorted(glob.glob(os.path.join(root, '*.mkv'))):
            r = osd.calibrate(f, osd.DEFAULT_RECT)
            prof = core.Profile()
            evs = [e for e in core.number_events(core.process([f], [r['offset']], prof)) if e['end'] - e['start'] >= 0.6]
            out = _mkdtemp()
            core.save_frames_and_clips(evs, prof, out, files=[f], bases=[r['offset']])
            for e in evs:
                self.assertEqual(len(e['shots']), 3, (f, core.fmt_time(e['start'])))
                if e['valid']:
                    self.assertTrue(e['clip'] and os.path.getsize(os.path.join(out, '短片', e['clip'])) > 0, (f, core.fmt_time(e['start'])))
                else:
                    self.assertEqual(e['clip'], '')
                    self.assertTrue(any('變化最大' in x for x in e['shots']), e['shots'])
                img = cv2.imdecode(np.fromfile(os.path.join(out, '截圖', e['shots'][0]), np.uint8), cv2.IMREAD_COLOR)
                d, txt, _w = osd.read_datetime(img, osd.DEFAULT_RECT, osd.DEFAULT_FORMAT, T)
                self.assertIsNotNone(d, (f, txt))
                # 畫面上的秒數是「無條件捨去」顯示：車頭時間 28.0 秒時字幕可能還是 27（差 0.1 秒內的對齊誤差）
                lag = e['start'] - d.timestamp()
                self.assertTrue(-0.2 <= lag <= 1.2, (f, txt, core.fmt_time(e['start'])))
            fr, pos = core.read_frame_at(f, 550.0)        # 預覽：影片後段也要讀得到
            self.assertIsNotNone(fr, f)
            self.assertLess(abs(pos - 550.0), 0.2, f)


if __name__ == '__main__':
    unittest.main()


# ---------------------------------------------------------------- 打包時只跑快的測試（#72）
# GitHub Actions（環境變數 CI=true）打包時略過下面這些花時間長的測試（每項 18 秒以上，合計約 20 分鐘），
# 讓打包回到幾分鐘完成；Claude 每次交付前在開發環境全部跑完（沒有 CI），交付說明會寫出結果。
# 要在 GitHub 上也全部跑：設定環境變數 TW_FULL=1。build.yml 不用改。
SLOW_TESTS = {
    'TestRound2': ['test46_static_limit_end_unknown', 'test47_static_limit_across_files', 'test51_vertical_sunlight',
                   'test53_ir_switch_no_train', 'test53b_ir_switch_then_train', 'test56_train_standing_at_start',
                   'test59_long_event_clip'],
    'TestRound3': ['test67a_stop_longer_than_max_static'],
    'TestRound4': ['test71a_idle_camera_moved_stops', 'test71b_camera_moved_during_train_stops', 'test71c_no_false_camera_stop'],
    'TestRound5': ['test75a_move_at_30s_inside_initial_window', 'test75b_move_at_3s', 'test75c_train_at_start_no_move',
                   'test76d_gap_same_place_brightness', 'test76e_gap_camera_moved', 'test76f_gap_day_to_ir'],
    'TestRound6': ['test79a_crossing_a_longer', 'test79b_crossing_b_leaves_last', 'test79c_extreme_b_stops'],
    'TestRound8': ['test82a_profile_anchor_same_place', 'test82b_moved_before_batch', 'test83a_slow_drift',
                   'test83b_jitter_no_false_alarm', 'test83c_drift_and_back', 'test84a_rotation', 'test84b_zoom',
                   'test84c_ir_switch_no_stop', 'test84d_uncertain_then_recovered',
                   'test87_error_keeps_results'],
    'TestScenarios': ['test8_local_sunlight_no_train', 'test8b_train_during_sunlight'],
}
_SKIP_SLOW = os.environ.get('CI') == 'true' and not os.environ.get('TW_FULL')
for _cls, _names in SLOW_TESTS.items():
    for _n in _names:
        _c = globals()[_cls]
        setattr(_c, _n, unittest.skipIf(_SKIP_SLOW, '打包時略過（花時間長，交付前已在開發環境全部跑過）')(getattr(_c, _n)))

del _cls, _names, _n, _c

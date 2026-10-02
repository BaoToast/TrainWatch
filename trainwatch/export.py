"""結果輸出：Excel（列車通過紀錄／已排除／設定與影片）、CSV、results.json（可重新開啟檢視修改）"""
from __future__ import annotations

import csv
import datetime as dt
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .core import Profile, fmt_time

TYPES = ['', '自強', '莒光', '區間車（電車）', '普悠瑪', '太魯閣', '機動貨車（機貨車）', '貨運列車', '其他']
DIRS = ['往右', '往左', '往上', '往下', '無法判定']

KEEP = ('no', 'start', 'end', 'first_change', 'last_change', 'start_file', 'end_file', 'start_pos', 'end_pos',
        'coverage', 'both', 'peak', 'direction', 'dir_on', 'dir_off', 'dir_confident', 'night', 'valid', 'reason',
        'need_check', 'shots', 'clip', 'train_type', 'crossing', 'note', 'checked',
        'end_known', 'start_known', 'media_tag', 'peak_t', 'peak_file', 'peak_pos', 'start_shift', 'end_shift', 'check_items')


def clean(ev):
    d = {k: ev.get(k) for k in KEEP}
    for k in ('train_type', 'note', 'clip', 'reason'):
        d[k] = d.get(k) or ''
    for k in ('crossing', 'checked', 'need_check', 'valid', 'night', 'dir_confident'):
        d[k] = bool(d.get(k))
    d['shots'] = d.get('shots') or []
    d['end_known'] = ev.get('end_known') is not False          # v1.0.8 以前的結果沒有這個欄位＝已確認
    d['start_known'] = ev.get('start_known') is not False
    d['check_items'] = list(ev.get('check_items') or [])
    return d


def save_json(path, events, profile: Profile, files, timing):
    """timing：每個檔案的時間對照 [{'offset':…, 'source':'畫面字幕'/'手動', 'msg':…}]"""
    data = dict(app='列車通過判讀', version=2, saved=dt.datetime.now().isoformat(timespec='seconds'),
                profile=json.loads(profile.to_json()), files=list(files), timing=list(timing),
                events=[clean(e) for e in events])
    tmp = path + '.tmp'
    with open(tmp, 'w', encoding='utf-8') as fh:
        json.dump(data, fh, ensure_ascii=False, indent=1)
    os.replace(tmp, path)


def load_json(path):
    with open(path, encoding='utf-8') as fh:
        data = json.load(fh)
    prof = Profile.from_dict(data.get('profile'))
    return data.get('events', []), prof, data.get('files', []), data.get('timing', [])


def fmt_pos(p):
    if p is None:
        return ''
    p = float(p)
    return '%d:%02d.%d' % (int(p // 60), int(p % 60), int((p - int(p)) * 10))


def row_values(ev, valid=True):
    d = dt.datetime.fromtimestamp(ev['start'])
    need = '是' if ev.get('need_check') and not ev.get('checked') else ''
    base = [ev['no'] if valid else 'X%d' % ev['no'], d.strftime('%Y-%m-%d'), fmt_time(ev['start']), fmt_time(ev['end']),
            round(ev['end'] - ev['start'], 1), ev.get('direction', '')]
    if valid:
        base += [ev.get('train_type', ''), '是' if ev.get('crossing') else '', need, '已確認' if ev.get('checked') else '',
                 ev.get('reason', ''), ev.get('note', '')]
    else:
        base += [ev.get('reason', ''), ev.get('note', '')]
    base += ['、'.join(ev.get('shots') or []), ev.get('clip', ''), os.path.basename(ev.get('start_file') or ''),
             fmt_pos(ev.get('start_pos'))]
    return base


HEAD_VALID = ['編號', '日期', '車頭到達', '車尾離開', '通過秒數', '方向', '車種', '雙向交會', '需人工確認', '人工確認',
              '系統備註', '使用者備註', '截圖檔', '短片檔', '影片檔', '影片內時間']
HEAD_EXCL = ['編號', '日期', '開始', '結束', '秒數', '方向', '排除原因', '使用者備註', '截圖檔', '短片檔', '影片檔', '影片內時間']


def _textw(s):
    return sum(2 if ord(ch) > 0x2E7F else 1 for ch in str(s))


def _sheet(ws, head, rows, wrap_cols=(), wrap_width=50):
    thin = Side(style='thin', color='999999')
    bd = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws.append(head)
    for c in range(1, len(head) + 1):
        cell = ws.cell(1, c)
        cell.font = Font(bold=True)
        cell.fill = PatternFill('solid', fgColor='DDEBF7')
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = bd
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.border = bd
            wrap = cell.column in wrap_cols
            cell.alignment = Alignment(vertical='center', wrap_text=wrap, horizontal='left' if wrap else 'center')
    for c in range(1, len(head) + 1):          # 欄寬依內容（中文字算 2），長欄位設上限並換行
        w = max([_textw(ws.cell(r, c).value) for r in range(1, ws.max_row + 1) if ws.cell(r, c).value is not None] or [4])
        if c in wrap_cols:
            w = min(w, wrap_width)
        ws.column_dimensions[get_column_letter(c)].width = w + 2
    ws.freeze_panes = 'A2'
    if ws.max_row >= 1:
        ws.auto_filter.ref = ws.dimensions


def write_excel(path, events, profile: Profile, files, timing):
    wb = Workbook()
    ws = wb.active
    ws.title = '列車通過紀錄'
    valid = [e for e in events if e.get('valid')]
    excl = [e for e in events if not e.get('valid')]
    _sheet(ws, HEAD_VALID, [row_values(e) for e in valid], wrap_cols=(11, 12, 13))
    yellow = PatternFill('solid', fgColor='FFF2CC')       # 需人工確認的列標淡黃色
    for i, e in enumerate(valid, start=2):
        if e.get('need_check') and not e.get('checked'):
            for c in range(1, len(HEAD_VALID) + 1):
                ws.cell(i, c).fill = yellow
    ws2 = wb.create_sheet('已排除')
    _sheet(ws2, HEAD_EXCL, [row_values(e, False) for e in excl], wrap_cols=(7, 8, 9))

    ws3 = wb.create_sheet('設定與影片')
    rows = [['項目', '內容'],
            ['監測站', profile.name],
            ['參考線（x1,y1,x2,y2）', str(sum(profile.ref_line, []))],
            ['參考線寬度', profile.ref_width],
            ['軌道範圍（x0,y0,x1,y1）', str(list(profile.track_rect))],
            ['時間字幕位置（x0,y0,x1,y1）', str(list(profile.osd_rect))],
            ['時間字幕格式', profile.osd_format],
            ['佔用門檻（開始／結束）', '%s／%s' % (profile.on_level, profile.off_level)],
            ['合併間隔（秒）', profile.merge_gap],
            ['最短秒數（少於排除）', profile.min_duration],
            ['最小軌道覆蓋率（低於排除）', profile.min_coverage],
            ['參考線兩側同時被擋的最小比例（低於排除）', profile.min_both],
            ['有效筆數', len(valid)],
            ['需人工確認筆數', sum(1 for e in valid if e.get('need_check') and not e.get('checked'))],
            ['已排除筆數', len(excl)],
            ['產生時間', dt.datetime.now().strftime('%Y-%m-%d %H:%M:%S')]]
    stops = [tm.get('stop') for tm in timing if tm.get('stop')]
    if stops:                                   # 攝影機位置改變而停止判讀（#71）
        rows.insert(1, ['⚠ 判讀中途停止', stops[0].replace('\n', '')])
    for r in rows:
        ws3.append(r)
    ws3.append([])
    ws3.append(['影片檔', '影片 0 秒的畫面時間', '時間來源', '畫面時間一致率', '說明'])
    hdr = ws3.max_row
    for f, tm in zip(files, timing):
        off = tm.get('offset')
        ws3.append([os.path.basename(f), fmt_time(off, True) if off is not None else '',
                    tm.get('source', ''), ('%d%%' % round(tm['support'] * 100)) if tm.get('support') is not None else '',
                    '；'.join(x for x in (tm.get('msg', ''), tm.get('stop', '').replace('\n', '')) if x)])
    for c, w in zip('ABCDE', (40, 26, 12, 14, 50)):
        ws3.column_dimensions[c].width = w
    for r in (1, hdr):
        for c in range(1, 6):
            ws3.cell(r, c).font = Font(bold=True)
    for r in range(1, ws3.max_row + 1):
        ws3.cell(r, 2).alignment = Alignment(wrap_text=True, vertical='center')
        ws3.cell(r, 5).alignment = Alignment(wrap_text=True, vertical='center')
    tmp = os.path.join(os.path.dirname(path), '_tmp_out.xlsx')
    wb.save(tmp)
    os.replace(tmp, path)


def write_csv(path, events):
    valid = [e for e in events if e.get('valid')]
    tmp = path + '.tmp'
    with open(tmp, 'w', encoding='utf-8-sig', newline='') as fh:   # utf-8-sig：Excel 直接開不會亂碼
        w = csv.writer(fh)
        w.writerow(HEAD_VALID)
        for e in valid:
            w.writerow(row_values(e))
    os.replace(tmp, path)


# ---------------------------------------------------------------- 噪音分析程式用
NOISE_SUFFIX = '10,FR,,A'        # 噪音分析程式固定要的欄位（使用者提供的範例格式）


def noise_lines(events):
    """每一筆列車一行：日期,進入時間,離開時間,10,FR,,A（例：2023/09/18,11:01:26,11:01:28,10,FR,,A）
    只輸出「是列車」的事件（已排除、人工改為非列車的都不輸出），依進入時間排序。
    秒數取整：進入時間無條件捨去、離開時間無條件進位，噪音時段一定包住整列車。"""
    import math
    out = []
    for e in sorted((e for e in events if e.get('valid')), key=lambda e: e['start']):
        a = dt.datetime.fromtimestamp(math.floor(e['start'] + 1e-6))
        b = dt.datetime.fromtimestamp(math.ceil(e['end'] - 1e-6))
        out.append('%s,%s,%s,%s' % (a.strftime('%Y/%m/%d'), a.strftime('%H:%M:%S'), b.strftime('%H:%M:%S'), NOISE_SUFFIX))
    return out


def write_noise_txt(path, events):
    """記事本格式（純英數字、Windows 換行），回傳寫了幾行"""
    lines = noise_lines(events)
    tmp = path + '.tmp'
    with open(tmp, 'w', encoding='ascii', newline='\r\n') as fh:
        for ln in lines:
            fh.write(ln + '\n')
    os.replace(tmp, path)
    return len(lines)

